"""Writing the workbooks we hand back."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from shared.spreadsheet import SheetColumn, write_workbook

COLUMNS = (SheetColumn(heading="Участник"), SheetColumn(heading="Голосов"))


def sheet_of(*rows: list[object]):
    payload = write_workbook(columns=COLUMNS, rows=rows, sheet_title="Итоги")
    book = load_workbook(io.BytesIO(payload.read()))
    return book.worksheets[0]


class TestWhatComesOut:
    def test_the_heading_row_comes_first(self) -> None:
        sheet = sheet_of(["Хор", 12])

        assert [cell.value for cell in sheet[1]] == ["Участник", "Голосов"]

    def test_values_land_in_their_columns(self) -> None:
        sheet = sheet_of(["Хор", 12], ["Дуэт", 7])

        assert [cell.value for cell in sheet[2]] == ["Хор", 12]
        assert [cell.value for cell in sheet[3]] == ["Дуэт", 7]

    def test_numbers_stay_numbers_so_the_recipient_can_sum_them(self) -> None:
        sheet = sheet_of(["Хор", 12])

        assert sheet.cell(row=2, column=2).value == 12

    def test_the_header_stays_on_screen_while_a_long_list_scrolls(self) -> None:
        assert sheet_of(["Хор", 1]).freeze_panes == "A2"

    def test_the_sheet_keeps_the_name_it_was_given(self) -> None:
        assert sheet_of(["Хор", 1]).title == "Итоги"

    def test_a_sheet_name_excel_would_refuse_is_cleaned_rather_than_fatal(self) -> None:
        payload = write_workbook(columns=COLUMNS, rows=[], sheet_title="Итоги: 2026/2027")
        book = load_workbook(io.BytesIO(payload.read()))

        assert book.worksheets[0].title == "Итоги 20262027"


class TestFormulasDoNotLeaveTheBuilding:
    @pytest.mark.parametrize("dangerous", ["=1+1", "+SUM(A1)", "-2+3", "@SUM(A1)"])
    def test_a_value_excel_would_execute_is_written_as_text(self, dangerous: str) -> None:
        """An export is a report. A report must not run on the recipient's machine."""
        sheet = sheet_of([dangerous, 1])

        assert sheet.cell(row=2, column=1).value == f"'{dangerous}"

    def test_an_ordinary_name_is_left_exactly_as_it_was(self) -> None:
        sheet = sheet_of(["Ансамбль «Ритм-2»", 1])

        assert sheet.cell(row=2, column=1).value == "Ансамбль «Ритм-2»"


class TestLargeExports:
    def test_it_writes_from_a_generator_without_holding_the_whole_list(self) -> None:
        rows = ([f"Участник {index}", index] for index in range(1000))

        payload = write_workbook(columns=COLUMNS, rows=rows)
        sheet = load_workbook(io.BytesIO(payload.read())).worksheets[0]

        assert sheet.max_row == 1001
        assert sheet.cell(row=1001, column=1).value == "Участник 999"
