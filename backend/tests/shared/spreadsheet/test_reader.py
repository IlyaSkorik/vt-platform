"""Reading the spreadsheets organizers actually send."""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook

from shared.spreadsheet import MAX_ROWS, Column, SheetReadError, read_rows

COLUMNS = (
    Column(key="number", names=("Номер", "№")),
    Column(key="title", names=("Участник", "ФИО", "Коллектив"), required=True),
    Column(key="city", names=("Город",)),
)


def workbook_of(*rows: list[object]) -> io.BytesIO:
    """Build an in-memory .xlsx whose sheet holds exactly ``rows``."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    payload = io.BytesIO()
    book.save(payload)
    payload.seek(0)
    return payload


class TestTheHappyPath:
    def test_it_reads_a_plain_table(self) -> None:
        source = workbook_of(["Номер", "Участник", "Город"], [1, "Ансамбль «Ритм»", "Минск"])

        result = read_rows(source, columns=COLUMNS)

        assert result.ok
        assert result.rows == ({"number": "1", "title": "Ансамбль «Ритм»", "city": "Минск"},)

    def test_a_number_typed_as_a_number_does_not_grow_a_decimal_tail(self) -> None:
        source = workbook_of(["Номер", "Участник"], [12, "Хор"])

        assert read_rows(source, columns=COLUMNS).rows[0]["number"] == "12"


class TestWhatRealFilesLookLike:
    def test_the_columns_may_be_in_any_order(self) -> None:
        source = workbook_of(["Город", "Участник", "Номер"], ["Брест", "Дуэт", 4])

        row = read_rows(source, columns=COLUMNS).rows[0]

        assert row == {"city": "Брест", "title": "Дуэт", "number": "4"}

    def test_a_heading_is_matched_through_its_synonyms(self) -> None:
        source = workbook_of(["№", "ФИО"], [1, "Иванова Мария"])

        assert read_rows(source, columns=COLUMNS).rows[0]["title"] == "Иванова Мария"

    def test_case_and_stray_spaces_in_a_heading_do_not_matter(self) -> None:
        source = workbook_of(["  НОМЕР ", "участник"], [1, "Соло"])

        assert read_rows(source, columns=COLUMNS).rows[0]["title"] == "Соло"

    def test_a_title_above_the_table_does_not_hide_the_header(self) -> None:
        source = workbook_of(
            ["Список номеров фестиваля «Весна»"],
            [],
            ["Номер", "Участник", "Город"],
            [1, "Трио", "Гомель"],
        )

        result = read_rows(source, columns=COLUMNS)

        assert result.ok
        assert result.rows[0]["title"] == "Трио"

    def test_trailing_empty_rows_are_not_errors(self) -> None:
        source = workbook_of(["Номер", "Участник"], [1, "Квартет"], [None, None], [None, None])

        result = read_rows(source, columns=COLUMNS)

        assert result.ok
        assert len(result.rows) == 1

    def test_a_column_the_file_lacks_is_reported_rather_than_dropped_silently(self) -> None:
        source = workbook_of(["Номер", "Участник"], [1, "Дуэт"])

        assert read_rows(source, columns=COLUMNS).missing_columns == ("Город",)


class TestWhatGoesWrong:
    def test_a_row_missing_a_required_value_is_reported_with_its_excel_row_number(self) -> None:
        source = workbook_of(["Номер", "Участник"], [1, "Хор"], [2, ""], [3, "Соло"])

        result = read_rows(source, columns=COLUMNS)

        assert not result.ok
        assert [error.row for error in result.errors] == [3]
        assert "«Участник»" in result.errors[0].message
        # The rows that were fine still come through — one bad line is not a bad file.
        assert len(result.rows) == 2

    def test_a_file_without_the_required_column_is_refused_whole(self) -> None:
        source = workbook_of(["Номер", "Город"], [1, "Минск"])

        with pytest.raises(SheetReadError, match="обязательных колонок"):
            read_rows(source, columns=COLUMNS)

    def test_a_file_with_no_recognizable_header_is_refused(self) -> None:
        source = workbook_of(["овощи", "фрукты"], ["морковь", "яблоко"])

        with pytest.raises(SheetReadError, match="заголовк"):
            read_rows(source, columns=COLUMNS)

    def test_something_that_is_not_a_workbook_is_refused_by_its_content(self) -> None:
        """The extension and the content type are the caller's word, not evidence."""
        source = io.BytesIO(b"\xd0\xad\xd1\x82\xd0\xbe \xd0\xbd\xd0\xb5 excel")

        with pytest.raises(SheetReadError, match="не файл Excel"):
            read_rows(source, columns=COLUMNS)

    def test_an_empty_file_is_refused(self) -> None:
        with pytest.raises(SheetReadError, match="пустой"):
            read_rows(io.BytesIO(b""), columns=COLUMNS)

    def test_a_zip_that_is_not_a_workbook_is_refused(self) -> None:
        payload = io.BytesIO()
        with zipfile.ZipFile(payload, "w") as archive:
            archive.writestr("hello.txt", "not a workbook")
        payload.seek(0)

        with pytest.raises(SheetReadError, match="прочитать файл Excel"):
            read_rows(payload, columns=COLUMNS)

    def test_a_file_that_runs_past_the_row_ceiling_is_refused(self) -> None:
        book = Workbook(write_only=True)
        sheet = book.create_sheet()
        sheet.append(["Номер", "Участник"])
        for index in range(MAX_ROWS + 5):
            sheet.append([index, f"Участник {index}"])
        payload = io.BytesIO()
        book.save(payload)
        payload.seek(0)

        with pytest.raises(SheetReadError, match="строк"):
            read_rows(payload, columns=COLUMNS)
