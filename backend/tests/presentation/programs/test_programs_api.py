"""Integration tests for the organizer Programme constructor REST edge.

Driven through DRF's ``APIClient`` as organization staff against real PostgreSQL: what
only exists at the edge is the staff gate, the wire shape, and the fact that acts come
from the Applications context rather than from a fixture of this one.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from typing import Any
from unittest.mock import patch

import pytest
from django.core import mail
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.applications.infrastructure.models.base import uuid7
from apps.applications.infrastructure.models.submission import SubmissionModel
from apps.competitions.infrastructure.models.competition import CompetitionModel
from apps.competitions.infrastructure.models.taxonomy import CategoryModel
from apps.identity.infrastructure.models.user import UserModel
from apps.identity.presentation.api.v1.auth import composition as identity_composition
from apps.organization.application.dto.commands import ProvisionOwnerCommand
from apps.organization.composition import build_provision_owner_handler
from apps.programs.application.commands import UpdateProgramCommand
from apps.programs.composition import build_program_service
from apps.programs.domain.entities import Program
from apps.programs.domain.exceptions import ProgramConcurrencyError
from apps.programs.infrastructure.repositories import DjangoProgramRepository
from apps.teachers.infrastructure.models import TeacherModel

pytestmark = pytest.mark.django_db

_PASSWORD = "V3ry$ecurePass!"
_PROGRAMS = "/api/v1/staff/programs"
TEN_AM = datetime(2026, 8, 1, 10, 0, tzinfo=UTC)


def _authenticated(client: APIClient, email: str) -> str:
    registration = client.post(
        "/api/v1/auth/registration",
        {"email": email, "password": _PASSWORD, "firstName": "Ольга", "lastName": "Петрова"},
        format="json",
    )
    access = str(identity_composition.build_token_service().issue(subject=registration.data["data"]["id"]).access_token)
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {access}")
    return access


@pytest.fixture
def staff_client() -> APIClient:
    """An organizer: authenticated *and* active organization staff."""
    client = APIClient()
    access = _authenticated(client, "organizer@example.com")
    build_provision_owner_handler().handle(ProvisionOwnerCommand(user_id=str(AccessToken(access)["user_id"])))  # type: ignore[arg-type]
    return client


@pytest.fixture
def outsider_client() -> APIClient:
    """Authenticated, but not staff."""
    client = APIClient()
    _authenticated(client, "participant@example.com")
    return client


@pytest.fixture
def scope() -> tuple[str, str]:
    """A published competition with one nomination; returns (competition_id, nomination_id)."""
    competition = CompetitionModel.objects.create(organizer_id=uuid.uuid4(), status="PUBLISHED", title="Весенний фестиваль", season="2025/2026")
    nomination = CategoryModel.objects.create(competition=competition, name="Вокал")
    return str(competition.id), str(nomination.id)


def _act(
    *,
    competition_id: str,
    nomination_id: str,
    title: str,
    seconds: int,
    performer: str = "Иванов Иван",
    comment: str = "",
    status: str = "ACCEPTED",
) -> str:
    """Seed one competing act in the Applications context."""
    submission = SubmissionModel.objects.create(
        competition_id=competition_id,
        # v7 и не иначе: ``UserId`` отказывает любой другой версии, а программа
        # теперь не только читает заявку, но и пишет в неё статус — через агрегат,
        # который этот идентификатор разбирает.
        applicant_id=uuid7(),
        status=status,
        performance={
            "title": title,
            "duration_seconds": seconds,
            "nomination_id": nomination_id,
            "discipline_id": str(uuid7()),
            "age_category_id": str(uuid7()),
            "participant_group_id": str(uuid7()),
            "participation_format_id": str(uuid7()),
            "comment": comment,
        },
        participants=[{"id": str(uuid7()), "display_name": performer, "member_count": 1}],
        # Полная форма, какую пишет ``SubmissionMapper``: программа теперь не только
        # читает заявку, но и грузит её агрегатом, чтобы проставить статус, — а он
        # разбирает каждое вложенное поле.
        teachers=[
            {
                "id": str(uuid7()),
                "first_name": "Мария",
                "last_name": "Иванова",
                "middle_name": None,
                "role": "COACH",
                "directory_teacher_id": None,
            }
        ],
    )
    return str(submission.id)


def _applicant(email: str, first_name: str, last_name: str) -> UserModel:
    """An account an announcement can actually be delivered to."""
    return UserModel.objects.create(
        email=email,
        first_name=first_name,
        last_name=last_name,
        locale="ru",
        timezone="Europe/Minsk",
        password_algorithm="argon2",
        password_hash="x",
    )


def _create_program(client: APIClient, scope: tuple[str, str], **overrides: Any) -> dict[str, Any]:
    competition_id, nomination_id = scope
    body: dict[str, Any] = {
        "competitionId": competition_id,
        "nominationId": nomination_id,
        "name": "Вокал, 1 отделение",
        "startsAt": TEN_AM.isoformat().replace("+00:00", "Z"),
        **overrides,
    }
    response = client.post(_PROGRAMS, body, format="json")
    assert response.status_code == 201, response.data
    # DRF types `Response.data` as Any; naming the shape here keeps that out of callers.
    created: dict[str, Any] = response.data["data"]
    return created


class TestAccess:
    def test_an_anonymous_caller_is_refused(self, scope: tuple[str, str]) -> None:
        assert APIClient().get(f"{_PROGRAMS}?competitionId={scope[0]}").status_code in (401, 403)

    def test_a_non_staff_user_is_refused(self, outsider_client: APIClient, scope: tuple[str, str]) -> None:
        assert outsider_client.get(f"{_PROGRAMS}?competitionId={scope[0]}").status_code == 403

    def test_listing_requires_a_competition(self, staff_client: APIClient) -> None:
        response = staff_client.get(_PROGRAMS)
        assert response.status_code == 400
        assert response.data["title"] == "Validation failed"


class TestProgramLifecycle:
    def test_a_program_is_created_and_listed(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        created = _create_program(staff_client, scope)
        assert created["name"] == "Вокал, 1 отделение"
        assert created["items"] == []

        listed = staff_client.get(f"{_PROGRAMS}?competitionId={scope[0]}")
        assert listed.status_code == 200
        assert [row["id"] for row in listed.data["data"]] == [created["id"]]
        assert listed.data["meta"]["total"] == 1

    def test_a_program_is_retrieved_and_renamed(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        renamed = staff_client.patch(f"{_PROGRAMS}/{program_id}", {"name": "Вокал, 2 отделение"}, format="json")
        assert renamed.status_code == 200
        assert renamed.data["data"]["name"] == "Вокал, 2 отделение"

    def test_a_blank_name_is_rejected_at_the_edge(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, _nomination_id = scope
        response = staff_client.post(_PROGRAMS, {"competitionId": competition_id, "name": ""}, format="json")
        assert response.status_code == 400

    def test_an_unknown_program_is_a_problem_document(self, staff_client: APIClient) -> None:
        response = staff_client.get(f"{_PROGRAMS}/{uuid.uuid4()}")
        assert response.status_code == 404
        assert response.data["title"] == "Resource not found"

    def test_a_program_can_be_deleted(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        assert staff_client.delete(f"{_PROGRAMS}/{program_id}").status_code == 204
        assert staff_client.get(f"{_PROGRAMS}/{program_id}").status_code == 404


class TestBuildingFromANomination:
    def test_candidates_come_from_the_applications_context(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=180)
        # Not competing yet: still under review.
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Черновик", seconds=180, status="UNDER_REVIEW")

        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.get(f"{_PROGRAMS}/{program_id}/candidates")
        assert response.status_code == 200
        assert sorted(row["title"] for row in response.data["data"]) == ["Калинка", "Катюша"]

    def test_filling_builds_the_running_order_with_computed_times(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт A", seconds=200)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт B", seconds=200)
        program_id = _create_program(staff_client, scope, bufferSeconds=40)["id"]

        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        assert filled.status_code == 200
        items = filled.data["data"]["items"]
        assert [item["schedule"]["position"] for item in items] == [1, 2]
        assert items[0]["schedule"]["startsAt"] == datetime(2026, 8, 1, 10, 0, tzinfo=UTC)
        assert items[1]["schedule"]["startsAt"] == datetime(2026, 8, 1, 10, 4, tzinfo=UTC)
        assert items[0]["schedule"]["slotSeconds"] == 240
        assert items[0]["durationSeconds"] == 200

    def test_a_single_act_is_added_by_id_and_carries_its_note(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(
            competition_id=competition_id,
            nomination_id=nomination_id,
            title="Катюша",
            seconds=200,
            performer="Иванов Иван",
            comment="Выход с точки 3",
        )
        program_id = _create_program(staff_client, scope)["id"]

        response = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")
        assert response.status_code == 201, response.data
        item = response.data["data"]["items"][0]
        assert item["kind"] == "PERFORMANCE"
        assert item["title"] == "Катюша"
        assert item["participants"] == ["Иванов Иван"]
        assert item["comment"] == "Выход с точки 3"
        assert item["teachers"] == ["Иванова Мария"]

    def test_an_act_that_is_not_competing_cannot_be_scheduled(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Черновик", seconds=180, status="UNDER_REVIEW")
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")
        assert response.status_code == 422

    def test_the_same_act_twice_is_a_conflict(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")
        again = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")
        assert again.status_code == 409


class TestBreaks:
    def test_a_break_is_inserted_and_re_times_the_rest(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт A", seconds=180)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт B", seconds=180)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")

        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/items",
            {"title": "Перерыв", "durationSeconds": 900, "position": 2},
            format="json",
        )
        assert response.status_code == 201, response.data
        items = response.data["data"]["items"]
        assert [item["kind"] for item in items] == ["PERFORMANCE", "BREAK", "PERFORMANCE"]
        assert items[1]["title"] == "Перерыв"
        assert items[2]["schedule"]["startsAt"] == datetime(2026, 8, 1, 10, 18, tzinfo=UTC)

    def test_a_break_is_edited_and_removed(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        created = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"title": "Обед", "durationSeconds": 1800}, format="json")
        item_id = created.data["data"]["items"][0]["id"]

        edited = staff_client.patch(f"{_PROGRAMS}/{program_id}/items/{item_id}", {"durationSeconds": 2400}, format="json")
        assert edited.status_code == 200
        assert edited.data["data"]["items"][0]["durationSeconds"] == 2400

        removed = staff_client.delete(f"{_PROGRAMS}/{program_id}/items/{item_id}")
        assert removed.status_code == 200
        assert removed.data["data"]["items"] == []

    def test_a_future_kind_of_service_block_already_works(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/items",
            {"title": "Награждение", "durationSeconds": 1800, "kind": "AWARD"},
            format="json",
        )
        assert response.status_code == 201, response.data
        assert response.data["data"]["items"][0]["kind"] == "AWARD"

    def test_an_unknown_kind_is_rejected(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/items",
            {"title": "Дискотека", "durationSeconds": 600, "kind": "PARTY"},
            format="json",
        )
        assert response.status_code == 400


class TestBufferAndOrder:
    def test_the_buffer_is_changed_and_the_schedule_follows(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт A", seconds=180)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт B", seconds=180)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")

        response = staff_client.patch(f"{_PROGRAMS}/{program_id}", {"bufferSeconds": 60}, format="json")
        assert response.status_code == 200
        assert response.data["data"]["bufferSeconds"] == 60
        # The act keeps its own length; only its slot grew.
        assert response.data["data"]["items"][0]["durationSeconds"] == 180
        assert response.data["data"]["items"][1]["schedule"]["startsAt"] == datetime(2026, 8, 1, 10, 4, tzinfo=UTC)

    def test_the_order_is_replaced_wholesale(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт A", seconds=100)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт B", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        ids = [item["id"] for item in filled.data["data"]["items"]]

        response = staff_client.put(f"{_PROGRAMS}/{program_id}/items/order", {"itemIds": list(reversed(ids))}, format="json")
        assert response.status_code == 200
        assert [item["id"] for item in response.data["data"]["items"]] == list(reversed(ids))

    def test_reordering_an_unpublished_programme_renumbers_it(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # Nothing has been handed out yet, so the numbers just follow the order.
        program_id, ids = self._three_act_program(staff_client, scope)

        response = staff_client.put(
            f"{_PROGRAMS}/{program_id}/items/order",
            {"itemIds": [ids[0], ids[2], ids[1]]},
            format="json",
        )
        assert response.status_code == 200
        items = response.data["data"]["items"]
        assert response.data["data"]["numbersFrozen"] is False
        assert [item["id"] for item in items] == [ids[0], ids[2], ids[1]]
        assert [item["number"] for item in items] == [1, 2, 3]

    def test_publishing_freezes_the_numbers(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id, _ids = self._three_act_program(staff_client, scope)
        created = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Экран"}, format="json")

        assert created.data["data"]["program"]["numbersFrozen"] is True
        assert created.data["data"]["program"]["numbersFrozenAt"] is not None

    def test_reordering_a_published_programme_keeps_the_numbers(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # A participant asked to perform later: the row moves, номер 3 stays номер 3, and
        # the programme reads 1, 3, 2 - which is what everybody was told.
        program_id, ids = self._three_act_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Экран"}, format="json")

        response = staff_client.put(
            f"{_PROGRAMS}/{program_id}/items/order",
            {"itemIds": [ids[0], ids[2], ids[1]]},
            format="json",
        )
        assert response.status_code == 200
        items = response.data["data"]["items"]
        assert [item["id"] for item in items] == [ids[0], ids[2], ids[1]]
        assert [item["number"] for item in items] == [1, 3, 2]
        # The running order itself is still contiguous - only the numbers jump.
        assert [item["schedule"]["position"] for item in items] == [1, 2, 3]

    def test_renumbering_is_a_separate_deliberate_action(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id, ids = self._three_act_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Экран"}, format="json")
        staff_client.put(f"{_PROGRAMS}/{program_id}/items/order", {"itemIds": [ids[2], ids[0], ids[1]]}, format="json")

        response = staff_client.post(f"{_PROGRAMS}/{program_id}/items/numbers")
        assert response.status_code == 200, response.data
        assert [item["number"] for item in response.data["data"]["items"]] == [1, 2, 3]

    @staticmethod
    def _three_act_program(staff_client: APIClient, scope: tuple[str, str]) -> tuple[str, list[str]]:
        competition_id, nomination_id = scope
        for index in range(3):
            _act(competition_id=competition_id, nomination_id=nomination_id, title=f"Акт {index}", seconds=120)
        program_id = _create_program(staff_client, scope)["id"]
        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        assert [item["number"] for item in filled.data["data"]["items"]] == [1, 2, 3]
        return program_id, [item["id"] for item in filled.data["data"]["items"]]

    def test_an_incomplete_order_is_rejected(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт A", seconds=100)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт B", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        first = filled.data["data"]["items"][0]["id"]

        response = staff_client.put(f"{_PROGRAMS}/{program_id}/items/order", {"itemIds": [first]}, format="json")
        assert response.status_code == 422

    def test_an_over_long_buffer_is_rejected(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        assert staff_client.patch(f"{_PROGRAMS}/{program_id}", {"bufferSeconds": 7200}, format="json").status_code == 400


class TestOneWriteDoesNotUndoAnother:
    """A program is read, changed and written back whole, by every endpoint there is.

    Which makes two overlapping writes a silent data loss: the later one restores every
    field to what *it* read, so the start time an organizer set a second earlier simply
    comes back empty after a reload. The version each copy was read at is what makes
    that refusable instead of invisible.
    """

    def test_a_start_set_on_an_untimed_programme_survives_a_reload(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        created = _create_program(staff_client, scope, startsAt=None)
        assert created["startsAt"] is None

        patched = staff_client.patch(f"{_PROGRAMS}/{created['id']}", {"startsAt": "2026-10-06T21:00:00.000Z"}, format="json")
        assert patched.status_code == 200, patched.data

        reloaded = staff_client.get(f"{_PROGRAMS}/{created['id']}")
        assert reloaded.data["data"]["startsAt"] == datetime(2026, 10, 6, 21, 0, tzinfo=UTC)

    def test_a_write_that_read_the_programme_too_early_is_refused(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope, startsAt=None)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")

        repository = DjangoProgramRepository()
        # One request has read the programme and not yet written it back.
        stale = repository.get(uuid.UUID(program_id))
        assert stale is not None and stale.starts_at is None

        # Another sets the start and commits in the meantime.
        build_program_service().update_program(UpdateProgramCommand(program_id=program_id, starts_at=TEN_AM))

        # The first now finishes its own work. It must not put the start back to None.
        stale.reorder([item.id for item in stale.items], now=TEN_AM)
        with pytest.raises(ProgramConcurrencyError):
            repository.save(stale)
        reloaded = repository.get(uuid.UUID(program_id))
        assert reloaded is not None and reloaded.starts_at == TEN_AM

    def test_a_settings_edit_replays_itself_rather_than_failing(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        """The settings are the one thing edited *while* the rest is being worked on.

        A partial edit means the same thing on a fresher copy — «поставь начало в 10:00»
        does not depend on what else moved — so it is replayed instead of reported.
        """
        program_id = _create_program(staff_client, scope, startsAt=None)["id"]
        collided = {"yet": False}
        read = DjangoProgramRepository.get

        def read_then_let_somebody_else_in(self: DjangoProgramRepository, pid: uuid.UUID) -> Program | None:
            program = read(self, pid)
            # Exactly once, between this read and its write, another write lands.
            if not collided["yet"]:
                collided["yet"] = True
                build_program_service().update_program(UpdateProgramCommand(program_id=str(pid), name="Переименована"))
            return program

        with patch.object(DjangoProgramRepository, "get", read_then_let_somebody_else_in):
            updated = build_program_service().update_program(UpdateProgramCommand(program_id=program_id, starts_at=TEN_AM))

        assert updated.starts_at == TEN_AM
        # And the write it collided with is still there: the retry re-read rather than overwrote.
        assert updated.name == "Переименована"


class TestTheProgrammeTable:
    def test_the_programme_exposes_every_column_it_can_carry(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program = _create_program(staff_client, scope)
        keys = [column["key"] for column in program["availableColumns"]]
        assert keys == [
            "POSITION",
            "TIME",
            "PARTICIPANT",
            "TITLE",
            "DURATION",
            "NOMINATION",
            "DISCIPLINE",
            "PARTICIPANT_GROUP",
            "PARTICIPANT_COUNT",
            "AGE_CATEGORY",
            "PARTICIPATION_FORMAT",
            "TEACHERS",
            "ORGANIZATION",
            "CONTACTS",
            "COMMENT",
        ]
        assert {column["key"]: column["label"] for column in program["availableColumns"]}["TEACHERS"] == "Педагог"

    def test_a_row_carries_a_value_for_every_column(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(
            competition_id=competition_id,
            nomination_id=nomination_id,
            title="Катюша",
            seconds=200,
            performer="Иванов Иван",
            comment="Выход с точки 3",
        )
        program_id = _create_program(staff_client, scope)["id"]
        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")

        item = filled.data["data"]["items"][0]
        assert item["title"] == "Катюша"
        assert item["participants"] == ["Иванов Иван"]
        assert item["teachers"] == ["Иванова Мария"]
        assert item["comment"] == "Выход с точки 3"
        # Every taxonomy column the table offers is resolved, not just the ids.
        for key in ("nomination", "discipline", "ageCategory", "participantGroup", "participationFormat"):
            assert key in item


class TestHowManyPeopleStandOnStage:
    """The headcount column, which is not the number of performer rows.

    A collective is one row in the entry and the headcount its applicant typed; a solo
    is one row and one person. The organizer needs the second number - it is how many
    chairs, microphones and diplomas the act costs - and ``len(participants)`` answers
    the wrong question for exactly the entries where it matters.
    """

    def test_a_collective_reports_the_headcount_it_stated(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=200)
        submission = SubmissionModel.objects.get(id=submission_id)
        submission.participants = [{"id": str(uuid.uuid4()), "display_name": "Ансамбль «Росинка»", "member_count": 14}]
        submission.save(update_fields=["participants"])

        program_id = _create_program(staff_client, scope)["id"]
        item = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]
        assert item["participantCount"] == 14
        assert len(item["participants"]) == 1

    def test_several_performers_without_a_headcount_are_counted_one_each(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Дуэт", seconds=200)
        submission = SubmissionModel.objects.get(id=submission_id)
        # No ``member_count`` at all: entries filled in before the field existed.
        submission.participants = [
            {"id": str(uuid.uuid4()), "display_name": "Иванов Иван"},
            {"id": str(uuid.uuid4()), "display_name": "Петрова Анна"},
        ]
        submission.save(update_fields=["participants"])

        program_id = _create_program(staff_client, scope)["id"]
        item = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]
        assert item["participantCount"] == 2

    def test_a_break_puts_nobody_on_stage(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        added = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"title": "Перерыв", "durationSeconds": 600}, format="json")
        assert added.status_code == 201, added.data
        assert added.data["data"]["items"][0]["participantCount"] == 0


class TestWhoToRingAboutAnAct:
    """The contacts column: the teacher's phone, and the applicant's when there is none."""

    def test_the_teacher_of_record_is_the_number_to_call(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        teacher = TeacherModel.objects.create(
            owner_id=uuid.uuid4(),
            first_name="Мария",
            last_name="Иванова",
            role="COACH",
            phone="+375291112233",
            status="ACTIVE",
        )
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        submission = SubmissionModel.objects.get(id=submission_id)
        submission.teachers = [
            {"first_name": "Мария", "last_name": "Иванова", "middle_name": None, "role": "COACH", "directory_teacher_id": str(teacher.id)}
        ]
        submission.save(update_fields=["teachers"])

        program_id = _create_program(staff_client, scope)["id"]
        item = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]
        assert item["contacts"] == ["+375291112233"]

    def test_an_act_with_no_reachable_teacher_falls_back_to_the_applicant(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        applicant = UserModel.objects.create(
            email="mother@example.com",
            phone="+375337778899",
            first_name="Анна",
            last_name="Петрова",
            locale="ru",
            timezone="Europe/Minsk",
            password_algorithm="argon2",
            password_hash="x",
        )
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        # The seeded teacher is hand-typed: a name on the entry and nowhere to ring.
        SubmissionModel.objects.filter(id=submission_id).update(applicant_id=applicant.id)

        program_id = _create_program(staff_client, scope)["id"]
        item = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]
        assert item["contacts"] == ["+375337778899"]

    def test_an_act_nobody_left_a_number_for_says_so_with_a_blank(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        item = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]
        assert item["contacts"] == []


class TestCustomColumns:
    def test_a_column_is_added_filled_and_published(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        item_id = filled.data["data"]["items"][0]["id"]

        created = staff_client.post(f"{_PROGRAMS}/{program_id}/columns", {"title": "Микрофон"}, format="json")
        assert created.status_code == 201, created.data
        column = created.data["data"]["availableColumns"][-1]
        assert column["label"] == "Микрофон"
        assert column["custom"] is True
        assert column["key"].startswith("custom:")

        written = staff_client.put(
            f"{_PROGRAMS}/{program_id}/items/{item_id}/cells",
            {"values": {column["key"]: "на стойке"}},
            format="json",
        )
        assert written.status_code == 200, written.data
        assert written.data["data"]["items"][0]["customValues"] == {column["key"]: "на стойке"}

        share = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {"name": "Звукорежиссёр", "columns": ["TIME", "TITLE", column["key"]]},
            format="json",
        )
        assert share.status_code == 201, share.data
        assert [c["label"] for c in share.data["data"]["share"]["columns"]] == ["Время", "Номер", "Микрофон"]

    def test_a_column_is_renamed_without_losing_its_values(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        item_id = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]["id"]
        key = staff_client.post(f"{_PROGRAMS}/{program_id}/columns", {"title": "Микрофон"}, format="json").data["data"]["availableColumns"][-1]["key"]
        staff_client.put(f"{_PROGRAMS}/{program_id}/items/{item_id}/cells", {"values": {key: "на стойке"}}, format="json")

        column_id = key.removeprefix("custom:")
        renamed = staff_client.patch(f"{_PROGRAMS}/{program_id}/columns/{column_id}", {"title": "Микрофон (тип)"}, format="json")
        assert renamed.status_code == 200, renamed.data
        assert renamed.data["data"]["availableColumns"][-1]["label"] == "Микрофон (тип)"
        assert renamed.data["data"]["items"][0]["customValues"] == {key: "на стойке"}

    def test_removing_a_column_clears_its_cells(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        item_id = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]["id"]
        key = staff_client.post(f"{_PROGRAMS}/{program_id}/columns", {"title": "Микрофон"}, format="json").data["data"]["availableColumns"][-1]["key"]
        staff_client.put(f"{_PROGRAMS}/{program_id}/items/{item_id}/cells", {"values": {key: "на стойке"}}, format="json")

        removed = staff_client.delete(f"{_PROGRAMS}/{program_id}/columns/{key.removeprefix('custom:')}")
        assert removed.status_code == 200, removed.data
        assert all(not column["custom"] for column in removed.data["data"]["availableColumns"])
        assert removed.data["data"]["items"][0]["customValues"] == {}

    def test_a_blank_column_title_is_rejected(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        assert staff_client.post(f"{_PROGRAMS}/{program_id}/columns", {"title": "  "}, format="json").status_code == 400

    def test_a_built_in_column_cannot_be_filled_in_by_hand(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        item_id = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"][0]["id"]

        response = staff_client.put(
            f"{_PROGRAMS}/{program_id}/items/{item_id}/cells",
            {"values": {"TITLE": "подменю название"}},
            format="json",
        )
        assert response.status_code == 400

    def test_publishing_an_unknown_custom_column_is_refused(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {"name": "Чужое", "columns": [f"custom:{uuid.uuid4()}"]},
            format="json",
        )
        assert response.status_code == 404


class TestPdfExport:
    """The organizer prints the programme, choosing which columns go on the page."""

    def _program_with_two_acts(self, staff_client: APIClient, scope: tuple[str, str]) -> str:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200, comment="Выход с точки 3")
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=200)
        program_id = str(_create_program(staff_client, scope)["id"])
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        return program_id

    def test_the_export_is_a_pdf_attachment(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)

        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.pdf?columns=POSITION,TIME,TITLE")
        assert response.status_code == 200, response.content[:200]
        assert response["Content-Type"] == "application/pdf"
        assert "attachment" in response["Content-Disposition"]
        # A real PDF, not an error page rendered with the wrong content type.
        assert response.content.startswith(b"%PDF-")
        assert len(response.content) > 1000

    def test_every_column_is_exported_when_none_is_chosen(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)
        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.pdf")
        assert response.status_code == 200
        assert response.content.startswith(b"%PDF-")

    def test_a_custom_column_can_be_exported(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)
        key = staff_client.post(f"{_PROGRAMS}/{program_id}/columns", {"title": "Микрофон"}, format="json").data["data"]["availableColumns"][-1]["key"]

        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.pdf?columns=TITLE,{key}")
        assert response.status_code == 200
        assert response.content.startswith(b"%PDF-")

    def test_an_empty_programme_still_exports(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # A programme is often printed before it is finished; an empty page beats an error.
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.pdf?columns=TIME,TITLE")
        assert response.status_code == 200
        assert response.content.startswith(b"%PDF-")

    def test_an_unknown_column_is_a_problem_document(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)
        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.pdf?columns=NOT_A_COLUMN")
        assert response.status_code == 400

    def test_an_unknown_programme_cannot_be_exported(self, staff_client: APIClient) -> None:
        assert staff_client.get(f"{_PROGRAMS}/{uuid.uuid4()}/export.pdf").status_code == 404

    def test_the_export_is_staff_only(self, outsider_client: APIClient, scope: tuple[str, str]) -> None:
        assert outsider_client.get(f"{_PROGRAMS}/{uuid.uuid4()}/export.pdf").status_code == 403


def _downloaded_rows(response: Any) -> list[list[str]]:
    """The streamed .xlsx as rows of text — what the recipient sees when they open it."""
    sheet = load_workbook(io.BytesIO(b"".join(response.streaming_content))).worksheets[0]
    return [["" if cell.value is None else str(cell.value) for cell in row] for row in sheet.iter_rows()]


class TestSpreadsheetExport:
    """The organizer takes the same table away in the format they keep working in."""

    def _program_with_two_acts(self, staff_client: APIClient, scope: tuple[str, str]) -> str:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=200)
        program_id = str(_create_program(staff_client, scope)["id"])
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        return program_id

    def test_the_export_is_one_flat_sheet_of_the_running_order(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)

        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.xlsx?columns=POSITION,TITLE")
        assert response.status_code == 200, response.content[:200]
        assert "spreadsheetml" in response["Content-Type"]
        assert "attachment" in response["Content-Disposition"]

        rows = _downloaded_rows(response)
        # Headings, then one row per slot — in running order, nothing merged or grouped.
        assert rows[0] == ["№", "Номер"]
        assert sorted(row[1] for row in rows[1:]) == ["Калинка", "Катюша"]

    def test_every_column_is_exported_when_none_is_chosen(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)
        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.xlsx")
        assert response.status_code == 200
        assert len(_downloaded_rows(response)[0]) > 2

    def test_an_empty_programme_still_exports(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.get(f"{_PROGRAMS}/{program_id}/export.xlsx?columns=TIME,TITLE")
        assert response.status_code == 200

    def test_an_unknown_column_is_a_problem_document(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = self._program_with_two_acts(staff_client, scope)
        assert staff_client.get(f"{_PROGRAMS}/{program_id}/export.xlsx?columns=NOT_A_COLUMN").status_code == 400

    def test_an_unknown_programme_cannot_be_exported(self, staff_client: APIClient) -> None:
        assert staff_client.get(f"{_PROGRAMS}/{uuid.uuid4()}/export.xlsx").status_code == 404

    def test_the_export_is_staff_only(self, outsider_client: APIClient, scope: tuple[str, str]) -> None:
        assert outsider_client.get(f"{_PROGRAMS}/{uuid.uuid4()}/export.xlsx").status_code == 403


class TestPublicEditions:
    def test_an_edition_is_published_with_its_columns_and_a_copyable_path(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Акт A", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        filled = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        item_id = filled.data["data"]["items"][0]["id"]

        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {
                "name": "Программа звукорежиссёра",
                "columns": ["TIME", "TITLE", "DURATION", "COMMENT"],
                "itemIds": [item_id],
            },
            format="json",
        )
        assert response.status_code == 201, response.data
        share = response.data["data"]["share"]
        assert share["name"] == "Программа звукорежиссёра"
        assert share["path"] == f"/program/{share['slug']}"
        assert [column["key"] for column in share["columns"]] == ["TIME", "TITLE", "DURATION", "COMMENT"]
        assert share["itemCount"] == 1
        assert share["coversWholeProgram"] is False

    def test_an_edition_published_without_columns_falls_back_to_the_defaults(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        created = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Экран у сцены"}, format="json")
        assert [column["key"] for column in created.data["data"]["share"]["columns"]] == [
            "POSITION",
            "TIME",
            "PARTICIPANT",
            "TITLE",
        ]

    def test_an_unknown_column_is_rejected(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {"name": "Нет", "columns": ["TIME", "SOMETHING"]},
            format="json",
        )
        assert response.status_code == 400

    def test_the_columns_of_a_live_link_can_be_changed(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        created = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {"name": "Ведущий", "columns": ["TIME", "TITLE"]},
            format="json",
        )
        share = created.data["data"]["share"]

        updated = staff_client.patch(
            f"{_PROGRAMS}/{program_id}/shares/{share['id']}",
            {"columns": ["TIME", "PARTICIPANT", "TITLE"]},
            format="json",
        )
        assert updated.status_code == 200, updated.data
        assert [column["key"] for column in updated.data["data"]["share"]["columns"]] == [
            "TIME",
            "PARTICIPANT",
            "TITLE",
        ]
        # The link itself is permanent, so anything already handed out keeps working.
        assert updated.data["data"]["share"]["slug"] == share["slug"]

    def test_an_empty_edition_update_is_rejected(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        created = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Ведущий"}, format="json")
        share_id = created.data["data"]["share"]["id"]
        assert staff_client.patch(f"{_PROGRAMS}/{program_id}/shares/{share_id}", {}, format="json").status_code == 400

    def test_editions_are_listed_and_can_be_revoked(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        created = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Экран у сцены"}, format="json")
        share_id = created.data["data"]["share"]["id"]

        listed = staff_client.get(f"{_PROGRAMS}/{program_id}/shares")
        assert listed.status_code == 200
        assert [row["id"] for row in listed.data["data"]] == [share_id]

        revoked = staff_client.delete(f"{_PROGRAMS}/{program_id}/shares/{share_id}")
        assert revoked.status_code == 200
        assert revoked.data["data"]["shares"] == []

    def test_an_edition_announces_itself_to_everyone_it_covers(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        """The gap this closes: the link exists and nobody knows it does."""
        competition_id, nomination_id = scope
        first = _applicant("petrova@example.com", "Анна", "Петрова")
        second = _applicant("sidorov@example.com", "Пётр", "Сидоров")
        katyusha = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        kalinka = _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=200)
        SubmissionModel.objects.filter(id=katyusha).update(applicant_id=first.id)
        SubmissionModel.objects.filter(id=kalinka).update(applicant_id=second.id)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        created = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Для участников"}, format="json")
        share = created.data["data"]["share"]
        mail.outbox.clear()

        announced = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares/{share['id']}/announce",
            {"note": "Сбор за час до начала."},
            format="json",
        )

        assert announced.status_code == 200, announced.data
        assert announced.data["data"] == {"notified": 2, "acts": 2, "skipped": 0}
        assert sorted(address for message in mail.outbox for address in message.to) == [
            "petrova@example.com",
            "sidorov@example.com",
        ]
        letter = next(message for message in mail.outbox if message.to == ["petrova@example.com"])
        assert letter.subject == "Программа опубликована"
        # The number, the time and the act — what the reader is actually being told —
        # and the link for everything else.
        assert "Катюша" in letter.body
        # Ten UTC read on the wall in the festival's own zone, the same minute the
        # constructor, the link and the PDF all show.
        assert "13:00" in letter.body
        assert "Сбор за час до начала." in letter.body
        assert share["path"] in letter.body
        # Their own act only: a letter listing the whole festival is a letter nobody reads.
        assert "Калинка" not in letter.body

    def test_a_studio_with_several_acts_is_written_to_once(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        studio = _applicant("studio@example.com", "Ольга", "Руководитель")
        for title in ("Катюша", "Калинка", "Барыня"):
            submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title=title, seconds=200)
            SubmissionModel.objects.filter(id=submission_id).update(applicant_id=studio.id)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        share = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Для участников"}, format="json").data["data"]["share"]
        mail.outbox.clear()

        announced = staff_client.post(f"{_PROGRAMS}/{program_id}/shares/{share['id']}/announce", {}, format="json")

        assert announced.data["data"] == {"notified": 1, "acts": 3, "skipped": 0}
        assert len(mail.outbox) == 1
        for title in ("Катюша", "Калинка", "Барыня"):
            assert title in mail.outbox[0].body

    def test_an_edition_cut_to_part_of_the_day_announces_only_that_part(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        first = _applicant("first@example.com", "Анна", "Петрова")
        second = _applicant("second@example.com", "Пётр", "Сидоров")
        katyusha = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        kalinka = _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=200)
        SubmissionModel.objects.filter(id=katyusha).update(applicant_id=first.id)
        SubmissionModel.objects.filter(id=kalinka).update(applicant_id=second.id)
        program_id = _create_program(staff_client, scope)["id"]
        items = staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill").data["data"]["items"]
        only_kalinka = next(item for item in items if item["title"] == "Калинка")
        share = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {"name": "Второе отделение", "itemIds": [only_kalinka["id"]]},
            format="json",
        ).data["data"]["share"]
        mail.outbox.clear()

        announced = staff_client.post(f"{_PROGRAMS}/{program_id}/shares/{share['id']}/announce", {}, format="json")

        assert announced.data["data"] == {"notified": 1, "acts": 1, "skipped": 0}
        assert [message.to for message in mail.outbox] == [["second@example.com"]]

    def test_an_act_nobody_can_be_written_to_about_is_counted_not_hidden(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        reachable = _applicant("reachable@example.com", "Анна", "Петрова")
        katyusha = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        # The second act's applicant has no account on the platform at all.
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=200)
        SubmissionModel.objects.filter(id=katyusha).update(applicant_id=reachable.id)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        share = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Для участников"}, format="json").data["data"]["share"]
        mail.outbox.clear()

        announced = staff_client.post(f"{_PROGRAMS}/{program_id}/shares/{share['id']}/announce", {}, format="json")

        # One letter sent and one act named as unreachable, rather than a failed request:
        # forty letters delivered beats none plus an error about the forty-first.
        assert announced.data["data"] == {"notified": 1, "acts": 1, "skipped": 1}
        assert [message.to for message in mail.outbox] == [["reachable@example.com"]]

    def test_breaks_are_not_announced_to_anybody(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        applicant = _applicant("solo@example.com", "Анна", "Петрова")
        katyusha = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        SubmissionModel.objects.filter(id=katyusha).update(applicant_id=applicant.id)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")
        staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"title": "Перерыв", "durationSeconds": 600}, format="json")
        share = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": "Для участников"}, format="json").data["data"]["share"]
        mail.outbox.clear()

        announced = staff_client.post(f"{_PROGRAMS}/{program_id}/shares/{share['id']}/announce", {}, format="json")

        assert announced.data["data"] == {"notified": 1, "acts": 1, "skipped": 0}
        assert "Перерыв" not in mail.outbox[0].body

    def test_announcing_a_link_from_another_programme_is_refused(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        other_id = _create_program(staff_client, scope, name="Другая")["id"]
        share = staff_client.post(f"{_PROGRAMS}/{other_id}/shares", {"name": "Чужая"}, format="json").data["data"]["share"]

        refused = staff_client.post(f"{_PROGRAMS}/{program_id}/shares/{share['id']}/announce", {}, format="json")

        assert refused.status_code == 404

    def test_an_outsider_cannot_announce(self, outsider_client: APIClient, scope: tuple[str, str]) -> None:
        assert outsider_client.post(f"{_PROGRAMS}/{uuid.uuid4()}/shares/{uuid.uuid4()}/announce", {}, format="json").status_code == 403

    def test_an_edition_cannot_select_a_foreign_item(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        response = staff_client.post(
            f"{_PROGRAMS}/{program_id}/shares",
            {"name": "Чужое", "itemIds": [str(uuid.uuid4())]},
            format="json",
        )
        assert response.status_code == 422

    def test_the_number_of_editions_is_not_limited(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program_id = _create_program(staff_client, scope)["id"]
        slugs = set()
        for name in ("Звукорежиссёр", "Ведущий", "Жюри", "Экран", "Утреннее отделение"):
            created = staff_client.post(f"{_PROGRAMS}/{program_id}/shares", {"name": name}, format="json")
            assert created.status_code == 201
            slugs.add(created.data["data"]["share"]["slug"])
        assert len(slugs) == 5


class TestArrangeEndpoint:
    """Generating the whole running order from rules (VT-86)."""

    def test_rules_in_the_body_build_the_programme(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        for title in ("Калинка", "Катюша"):
            _act(competition_id=competition_id, nomination_id=nomination_id, title=title, seconds=180)
        program = _create_program(staff_client, scope)

        response = staff_client.post(
            f"{_PROGRAMS}/{program['id']}/items/arrange",
            {"grouping": [], "orderWithinGroup": "TITLE", "breakRule": "EVERY_COUNT", "breakAfterItems": 1},
            format="json",
        )

        assert response.status_code == 200, response.data
        items = response.data["data"]["items"]
        assert [item["title"] for item in items] == ["Калинка", "Перерыв", "Катюша"]
        assert [item["kind"] for item in items] == ["PERFORMANCE", "BREAK", "PERFORMANCE"]

    def test_the_rules_come_back_on_the_programme(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)

        staff_client.post(
            f"{_PROGRAMS}/{program['id']}/items/arrange",
            {"grouping": ["NOMINATION"], "repeatGap": 3, "awardTitle": "Награждение"},
            format="json",
        )
        reloaded = staff_client.get(f"{_PROGRAMS}/{program['id']}")

        arrangement = reloaded.data["data"]["arrangement"]
        assert arrangement["grouping"] == ["NOMINATION"]
        assert arrangement["repeatGap"] == 3
        assert arrangement["awardTitle"] == "Награждение"

    def test_an_empty_body_rebuilds_on_the_saved_rules(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # The "пересобрать" button: no form is shown, so the programme's own rules are
        # the request.
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {"openingTitle": "Открытие"}, format="json")

        _act(competition_id=competition_id, nomination_id=nomination_id, title="Калинка", seconds=180)
        response = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")

        assert response.status_code == 200, response.data
        assert [item["title"] for item in response.data["data"]["items"]] == ["Открытие", "Калинка", "Катюша"]

    def test_a_programme_never_generated_has_no_rules_on_the_wire(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program = _create_program(staff_client, scope)
        assert program["arrangement"] is None

    def test_an_unknown_rule_is_a_validation_problem(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        program = _create_program(staff_client, scope)

        response = staff_client.post(
            f"{_PROGRAMS}/{program['id']}/items/arrange",
            {"orderWithinGroup": "BY_THE_MOON"},
            format="json",
        )

        assert response.status_code == 400

    def test_a_published_programme_is_refused_with_a_conflict(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")
        staff_client.post(f"{_PROGRAMS}/{program['id']}/shares", {"name": "Для зала"}, format="json")

        response = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")

        assert response.status_code == 409
        assert response.data["type"].endswith("/conflict")
        assert "опубликована" in response.data["detail"]

    def test_an_outsider_cannot_generate_a_programme(
        self,
        staff_client: APIClient,
        outsider_client: APIClient,
        scope: tuple[str, str],
    ) -> None:
        program = _create_program(staff_client, scope)
        response = outsider_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")
        assert response.status_code == 403


class TestUndoArrangeEndpoint:
    """«Отменить» — the one step back a regeneration leaves behind (VT-86)."""

    def test_the_previous_order_comes_back_whole(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/fill", {}, format="json")
        by_hand = staff_client.post(
            f"{_PROGRAMS}/{program['id']}/items",
            {"kind": "BREAK", "title": "Обед", "durationSeconds": 900},
            format="json",
        )
        assert by_hand.status_code in (200, 201), by_hand.data
        before = staff_client.get(f"{_PROGRAMS}/{program['id']}").data["data"]["items"]

        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")
        response = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange/undo", {}, format="json")

        assert response.status_code == 200, response.data
        assert [item["id"] for item in response.data["data"]["items"]] == [item["id"] for item in before]

    def test_the_programme_says_whether_the_undo_is_on_offer(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)

        assert program["canUndoArrange"] is False

        arranged = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")
        assert arranged.data["data"]["canUndoArrange"] is True

        undone = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange/undo", {}, format="json")
        assert undone.data["data"]["canUndoArrange"] is False

    def test_a_stale_undo_button_is_a_conflict_not_a_silent_no_op(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # Two windows open, both showing «Отменить»; the second press must say so
        # rather than appear to work and change nothing.
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")
        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange/undo", {}, format="json")

        response = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange/undo", {}, format="json")

        assert response.status_code == 409, response.data
        assert response.data["type"].endswith("conflict")

    def test_a_published_programme_is_refused_the_undo(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=180)
        program = _create_program(staff_client, scope)
        staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange", {}, format="json")
        staff_client.post(f"{_PROGRAMS}/{program['id']}/shares", {"name": "Экран"}, format="json")

        response = staff_client.post(f"{_PROGRAMS}/{program['id']}/items/arrange/undo", {}, format="json")

        assert response.status_code == 409, response.data
        reloaded = staff_client.get(f"{_PROGRAMS}/{program['id']}")
        assert reloaded.data["data"]["canUndoArrange"] is False


class TestTheRunningOrderDrivesTheStatus:
    """`SCHEDULED` ставится, когда номер действительно попал в программу.

    Через настоящую БД, а не через фейки: здесь проверяется связка двух контекстов
    целиком — конструктор программы пишет в Applications через ACL, и только так
    видно, что статус доехал до строки заявки.
    """

    def test_adding_an_act_marks_it_scheduled(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]

        staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")

        assert SubmissionModel.objects.get(id=submission_id).status == "SCHEDULED"

    def test_an_act_nobody_placed_stays_merely_accepted(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # Собственно то, ради чего правка: одобренный номер больше не выглядит
        # назначенным только потому, что его одобрили.
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Смуглянка", seconds=180)
        _create_program(staff_client, scope)

        assert SubmissionModel.objects.get(id=submission_id).status == "ACCEPTED"

    def test_pulling_the_slot_puts_the_act_back_to_accepted(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        added = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")
        item_id = added.data["data"]["items"][0]["id"]

        staff_client.delete(f"{_PROGRAMS}/{program_id}/items/{item_id}")

        assert SubmissionModel.objects.get(id=submission_id).status == "ACCEPTED"

    def test_filling_the_scope_marks_everything_it_pulled_in(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        first = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        second = _act(competition_id=competition_id, nomination_id=nomination_id, title="Смуглянка", seconds=180)
        program_id = _create_program(staff_client, scope)["id"]

        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")

        assert SubmissionModel.objects.get(id=first).status == "SCHEDULED"
        assert SubmissionModel.objects.get(id=second).status == "SCHEDULED"

    def test_deleting_the_programme_releases_its_acts(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        program_id = _create_program(staff_client, scope)["id"]
        staff_client.post(f"{_PROGRAMS}/{program_id}/items/fill")

        assert staff_client.delete(f"{_PROGRAMS}/{program_id}").status_code == 204

        assert SubmissionModel.objects.get(id=submission_id).status == "ACCEPTED"

    def test_an_act_in_a_second_programme_keeps_its_status(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # Конкурс ведёт несколько программ. Номер, вынутый из одной, может стоять в
        # другой — снимать статус по одной программе значит соврать про вторую.
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200)
        first = _create_program(staff_client, scope, name="Вокал, 1 отделение")["id"]
        second = _create_program(staff_client, scope, name="Вокал, 2 отделение")["id"]
        staff_client.post(f"{_PROGRAMS}/{first}/items", {"submissionId": submission_id}, format="json")
        in_second = staff_client.post(f"{_PROGRAMS}/{second}/items", {"submissionId": submission_id}, format="json")
        item_id = in_second.data["data"]["items"][0]["id"]

        staff_client.delete(f"{_PROGRAMS}/{second}/items/{item_id}")

        assert SubmissionModel.objects.get(id=submission_id).status == "SCHEDULED"

    def test_an_already_performed_act_is_not_dragged_back(self, staff_client: APIClient, scope: tuple[str, str]) -> None:
        # COMPLETED — это состоявшееся выступление, а не место в расписании: вынутая
        # строка не должна его отменять.
        competition_id, nomination_id = scope
        submission_id = _act(competition_id=competition_id, nomination_id=nomination_id, title="Катюша", seconds=200, status="COMPLETED")
        program_id = _create_program(staff_client, scope)["id"]
        added = staff_client.post(f"{_PROGRAMS}/{program_id}/items", {"submissionId": submission_id}, format="json")
        item_id = added.data["data"]["items"][0]["id"]

        staff_client.delete(f"{_PROGRAMS}/{program_id}/items/{item_id}")

        assert SubmissionModel.objects.get(id=submission_id).status == "COMPLETED"
