"""Filling a ballot from a spreadsheet — the path a standalone festival takes."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

pytestmark = pytest.mark.django_db

_BALLOTS = "/api/v1/staff/ballots"


def workbook_of(*rows: list[object]) -> io.BytesIO:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    payload = io.BytesIO()
    book.save(payload)
    payload.seek(0)
    payload.name = "nominees.xlsx"
    return payload


def upload(client: APIClient, ballot_id: str, payload: io.BytesIO, *, dry_run: bool = False):
    query = "?dryRun=true" if dry_run else ""
    return client.post(
        f"{_BALLOTS}/{ballot_id}/nominees/import{query}",
        {"file": payload},
        format="multipart",
    )


class TestTheTemplate:
    def test_the_organizer_is_handed_a_file_to_fill_in(self, staff_client: APIClient, ballot: dict) -> None:
        """Without an example file the organizer sends their own shape and we lose a day."""
        response = staff_client.get(f"{_BALLOTS}/{ballot['id']}/nominees/template.xlsx")

        assert response.status_code == 200
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")

    def test_the_template_names_the_columns_the_import_reads(self, staff_client: APIClient, ballot: dict) -> None:
        response = staff_client.get(f"{_BALLOTS}/{ballot['id']}/nominees/template.xlsx")
        sheet = load_workbook(io.BytesIO(b"".join(response.streaming_content))).worksheets[0]

        assert [cell.value for cell in sheet[1]] == ["Номер", "Участник", "Произведение", "Город"]

    def test_the_template_round_trips_back_through_the_import(self, staff_client: APIClient, ballot: dict) -> None:
        """The file we hand out has to be one we accept."""
        template = staff_client.get(f"{_BALLOTS}/{ballot['id']}/nominees/template.xlsx")
        payload = io.BytesIO(b"".join(template.streaming_content))
        payload.name = "template.xlsx"

        response = upload(staff_client, ballot["id"], payload)

        assert response.status_code == 200
        assert len(response.data["data"]["added"]) == 1


class TestImporting:
    def test_a_filled_in_list_becomes_nominees(self, staff_client: APIClient, ballot: dict) -> None:
        payload = workbook_of(
            ["Номер", "Участник", "Произведение", "Город"],
            [1, "Ансамбль «Заря»", "«Калинка»", "Минск"],
            [2, "Дуэт «Весна»", "«Вальс»", "Брест"],
        )

        response = upload(staff_client, ballot["id"], payload)

        assert response.status_code == 200, response.data
        added = response.data["data"]["added"]
        assert [row["title"] for row in added] == ["Ансамбль «Заря»", "Дуэт «Весна»"]
        assert added[0]["city"] == "Минск"
        assert added[0]["number"] == 1

    def test_the_columns_may_arrive_in_any_order_and_under_a_synonym(self, staff_client: APIClient, ballot: dict) -> None:
        """A list exported from another system does not match our column order."""
        payload = workbook_of(["Город", "ФИО", "№"], ["Гомель", "Иванова Мария", 4])

        response = upload(staff_client, ballot["id"], payload)

        added = response.data["data"]["added"]
        assert added[0]["title"] == "Иванова Мария"
        assert added[0]["city"] == "Гомель"
        assert added[0]["number"] == 4

    def test_a_merged_title_above_the_table_does_not_hide_the_header(self, staff_client: APIClient, ballot: dict) -> None:
        payload = workbook_of(
            ["Список номеров фестиваля «Весна»"],
            [],
            ["Номер", "Участник"],
            [1, "Трио"],
        )

        response = upload(staff_client, ballot["id"], payload)

        assert [row["title"] for row in response.data["data"]["added"]] == ["Трио"]

    def test_an_imported_list_lands_after_what_is_already_there(self, staff_client: APIClient, ballot: dict) -> None:
        payload = workbook_of(["Участник"], ["Новенький"])

        upload(staff_client, ballot["id"], payload)
        detail = staff_client.get(f"{_BALLOTS}/{ballot['id']}").data["data"]

        assert [row["position"] for row in detail["nominees"]] == [1, 2, 3]
        assert detail["nominees"][-1]["title"] == "Новенький"


class TestThePreview:
    def test_a_dry_run_changes_nothing(self, staff_client: APIClient, ballot: dict) -> None:
        """Importing 150 rows of somebody else's formatting is undone by hand."""
        payload = workbook_of(["Участник"], ["Пробный"])

        response = upload(staff_client, ballot["id"], payload, dry_run=True)

        assert response.data["data"]["dryRun"] is True
        assert len(response.data["data"]["added"]) == 1
        detail = staff_client.get(f"{_BALLOTS}/{ballot['id']}").data["data"]
        assert len(detail["nominees"]) == 2

    def test_a_bad_row_is_reported_with_the_number_excel_shows(self, staff_client: APIClient, ballot: dict) -> None:
        payload = workbook_of(["Номер", "Участник"], [1, "Хор"], [2, ""], [3, "Соло"])

        response = upload(staff_client, ballot["id"], payload, dry_run=True)

        errors = response.data["data"]["errors"]
        assert [error["row"] for error in errors] == [3]
        assert "«Участник»" in errors[0]["message"]
        # The good rows still come through: one bad line is not a bad file.
        assert len(response.data["data"]["added"]) == 2

    def test_a_column_the_file_lacks_is_named_rather_than_dropped_in_silence(self, staff_client: APIClient, ballot: dict) -> None:
        payload = workbook_of(["Участник"], ["Соло"])

        response = upload(staff_client, ballot["id"], payload, dry_run=True)

        assert "Город" in response.data["data"]["missingColumns"]


class TestRefusingAFile:
    def test_something_that_is_not_a_workbook_is_refused_by_its_content(self, staff_client: APIClient, ballot: dict) -> None:
        payload = io.BytesIO("Это не excel".encode())
        payload.name = "nominees.xlsx"

        response = upload(staff_client, ballot["id"], payload)

        assert response.status_code == 400
        assert "не файл Excel" in response.data["detail"]

    def test_a_file_without_a_recognizable_header_is_refused(self, staff_client: APIClient, ballot: dict) -> None:
        payload = workbook_of(["овощи", "фрукты"], ["морковь", "яблоко"])

        response = upload(staff_client, ballot["id"], payload)

        assert response.status_code == 400
        assert "заголовк" in response.data["detail"]

    def test_sending_no_file_at_all_says_so(self, staff_client: APIClient, ballot: dict) -> None:
        response = staff_client.post(f"{_BALLOTS}/{ballot['id']}/nominees/import", {}, format="multipart")

        assert response.status_code == 400


class TestExportingTheResults:
    def test_the_results_come_out_ranked(self, staff_client: APIClient, open_ballot: dict) -> None:
        """A results file is opened to find out who won."""
        loser, winner = open_ballot["nominees"]
        for _ in range(2):
            APIClient().post(
                f"/api/v1/public/ballots/{open_ballot['slug']}/votes",
                {"nomineeIds": [winner["id"]]},
                format="json",
            )
        APIClient().post(
            f"/api/v1/public/ballots/{open_ballot['slug']}/votes",
            {"nomineeIds": [loser["id"]]},
            format="json",
        )

        response = staff_client.get(f"{_BALLOTS}/{open_ballot['id']}/results.xlsx")
        sheet = load_workbook(io.BytesIO(b"".join(response.streaming_content))).worksheets[0]

        assert [cell.value for cell in sheet[1]][:3] == ["Место", "Номер", "Участник"]
        assert sheet.cell(row=2, column=3).value == winner["title"]
        assert sheet.cell(row=2, column=6).value == 2
        assert sheet.cell(row=3, column=6).value == 1

    def test_the_qr_for_the_poster_is_a_png(self, staff_client: APIClient, ballot: dict) -> None:
        response = staff_client.get(f"{_BALLOTS}/{ballot['id']}/qr.png")

        assert response.status_code == 200
        assert response["Content-Type"] == "image/png"
        assert response.content[:8] == b"\x89PNG\r\n\x1a\n"
